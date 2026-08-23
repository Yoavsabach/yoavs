#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_excel.py — בונה את המודל הכלכלי המלא כקובץ Excel.

עקרון מנחה: **גיליון "הנחות יסוד" הוא המקור היחיד**. כל מספר בכל גיליון אחר הוא
נוסחה שמצביעה חזרה אליו. יזם שמשנה מחיר מכירה למ"ר בתא אחד רואה את הרווח, ה-IRR
והתזרים מתעדכנים — וזו כל התכלית של מודל, להבדיל מטבלת מספרים מודפסת.

חריג יחיד ומכוון: גיליון "ניתוחי רגישות". חישוב הרגישות מחייב הרצה מחדש של לולאת
המימון הרבעונית לכל תא במטריצה, ואת זה אי אפשר לבטא בנוסחת Excel אחת. הערכים שם
מחושבים ב-Python ומסומנים ככאלה בגיליון עצמו, עם הוראה להריץ מחדש אחרי שינוי הנחות.

שימוש:
    python3 build_excel.py project.json -o "מודל כלכלי.xlsx"
"""

from __future__ import annotations

import argparse
import os
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from model import (build_model, derive_flags, deviation_notes, load_project,  # noqa: E402
                   num, pct, rate_scenarios, sensitivity_grid, DISCLAIMER)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import CellIsRule
except ImportError:  # pragma: no cover
    print("חסרה תלות. התקן:  pip install openpyxl", file=sys.stderr)
    raise

# --- תקני העיצוב של המשרד -------------------------------------------------
# פורמט מטבע חשבונאי עם ₪ בתוך התא (לא כתחילית מנותקת), כך שעמודות מספרים
# מיושרות על נקודת העשרוני גם כשהערכים בסדרי גודל שונים.
NIS = '_ "₪" * #,##0.00_ ;_ "₪" * (#,##0.00)_ ;_ "₪" * "-"??_ ;_ @_ '
NIS0 = '_ "₪" * #,##0_ ;_ "₪" * (#,##0)_ ;_ "₪" * "-"??_ ;_ @_ '
PCT = '0.0%'
SQM = '#,##0.0 "מ""ר"'
QTY = '#,##0'
DATE_FMT = 'dd/mm/yyyy'

NAVY = "1F3864"
ACCENT = "2E5C8A"
LIGHT = "DCE6F1"
GREY = "F2F2F2"
GREEN = "C6EFCE"
YELLOW = "FFEB9C"
RED = "FFC7CE"

HDR_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor=LIGHT)
TOT_FILL = PatternFill("solid", fgColor=GREY)
HDR_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Arial", size=14, bold=True, color=NAVY)
BOLD = Font(name="Arial", size=11, bold=True)
BASE = Font(name="Arial", size=11)
SMALL = Font(name="Arial", size=9, color="666666")

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

ASSUMPTIONS_SHEET = "הנחות יסוד"


def q(sheet_name: str) -> str:
    """שם גיליון בתוך נוסחה. שמות עבריים עם רווח חייבים גרשיים."""
    return "'%s'" % sheet_name


class Builder:
    def __init__(self, project):
        self.p = project
        self.m = build_model(project)
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        # מפתח → כתובת מוחלטת בגיליון ההנחות, למשל "'הנחות יסוד'!$C$7"
        self.ref = {}

    # -- עזרי כתיבה --------------------------------------------------------

    @staticmethod
    def _sum(col, first, last):
        """סכום עמודה עם הגנה מפני טווח הפוך/ריק.

        כשסקציה יוצאת בלי שורות נתונים (למשל 'עלויות עקיפות' בפרויקט יזמי
        מקוצר) מתקבל ``first > last`` והנוסחה הנאיבית ``=SUM(D4:D3)`` היא טווח
        הפוך — חלק ממנועי הגיליון מחזירים עליה ``#NULL!`` שמזהם את כל התזרים
        במורד הזרם. במקרה כזה מחזירים 0 מספרי חד-משמעי במקום טווח (לא המחרוזת
        "0" שהייתה נכתבת כתא טקסט ולא כמספר)."""
        if last < first:
            return 0
        return "=SUM(%s%d:%s%d)" % (col, first, col, last)

    def sheet(self, title, widths):
        ws = self.wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True   # כל גיליון RTL — לא רק הראשון
        ws.sheet_view.showGridLines = False
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        return ws

    def title(self, ws, row, text, span):
        c = ws.cell(row=row, column=1, value=text)
        c.font = TITLE_FONT
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
        ws.row_dimensions[row].height = 24
        return row + 2

    def header(self, ws, row, labels):
        for i, label in enumerate(labels, start=1):
            c = ws.cell(row=row, column=i, value=label)
            c.font = HDR_FONT
            c.fill = HDR_FILL
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BOX
        ws.row_dimensions[row].height = 30
        if ws.freeze_panes is None:
            ws.freeze_panes = ws.cell(row=row + 1, column=1)
        return row + 1

    def line(self, ws, row, values, fmts=None, bold=False, fill=None):
        for i, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=i, value=v)
            c.font = BOLD if bold else BASE
            c.border = BOX
            if fill:
                c.fill = fill
            if fmts and i - 1 < len(fmts) and fmts[i - 1]:
                c.number_format = fmts[i - 1]
            if isinstance(v, str) and not v.startswith("="):
                c.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        return row + 1

    def note(self, ws, row, text, span=6):
        c = ws.cell(row=row, column=1, value=text)
        c.font = SMALL
        c.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
        ws.row_dimensions[row].height = 28
        return row + 2

    # -- 1. הנחות יסוד -----------------------------------------------------

    def build_assumptions(self):
        ws = self.sheet(ASSUMPTIONS_SHEET, [22, 46, 18, 14, 40, 34])
        r = self.title(ws, 1, "הנחות יסוד — %s" % self.m["meta"].get("project_name", ""), 6)
        r = self.note(ws, r,
                      "זהו המקור היחיד לכל מספר במודל. כל גיליון אחר מקושר לכאן בנוסחה חיה — "
                      "שינוי ערך כאן מעדכן את המודל כולו. שיעורי אחוז נמסרים ביחידות אחוז (6 = 6%).")
        r = self.header(ws, r, ["קטגוריה", "פרמטר", "ערך", "יחידה", "מקור / אסמכתה", "הערת סטייה"])

        devs = {d["label"]: d for d in deviation_notes(self.p)}
        p = self.p

        def add(cat, key, label, value, unit, source="", fmt=None):
            nonlocal r
            dev = devs.get(label)
            self.line(ws, r, [cat, label, value, unit, source or "—",
                              dev["text"] if dev else ""],
                      fmts=[None, None, fmt, None, None, None])
            if dev:
                ws.cell(row=r, column=6).fill = PatternFill("solid", fgColor=YELLOW)
                ws.cell(row=r, column=3).fill = PatternFill("solid", fgColor=YELLOW)
            self.ref[key] = "%s!$C$%d" % (q(ASSUMPTIONS_SHEET), r)
            r += 1

        meta = p.get("meta", {})
        add("זיהוי", "gush", "גוש", meta.get("gush", ""), "", meta.get("plan", ""))
        add("זיהוי", "helka", "חלקה", meta.get("helka", ""), "")
        add("זיהוי", "migrash", "מגרש", meta.get("migrash", ""), "")

        land = p.get("land", {}) or {}
        src = land.get("source", "")
        add("קרקע", "land_price", "מחיר רכישת הקרקע", num(land.get("purchase_price")), "₪", src, NIS0)
        add("קרקע", "purchase_tax_rate", "שיעור מס רכישה", num(land.get("purchase_tax_rate"), 6), "%", "tax-rules.md")
        # המקור שייך לשורת מחיר הקרקע. הדבקתו על חמש שורות יוצרת מראה של
        # אסמכתה גם לשורות שאיש לא אימת.
        add("קרקע", "legal_pct", "שכ\"ט עו\"ד ועלויות עסקה", num(land.get("legal_pct")), "%",
            land.get("legal_source", ""))
        add("קרקע", "brokerage_pct", "דמי תיווך", num(land.get("brokerage_pct")), "%",
            land.get("brokerage_source", ""))
        add("קרקע", "dev_levies", "היטלי פיתוח בגין הקרקע", num(land.get("development_levies")), "₪",
            land.get("levies_source", ""), NIS0)
        add("קרקע", "land_other", "עלויות קרקע אחרות", num(land.get("other")), "₪",
            land.get("other_source", ""), NIS0)

        idx = p.get("indexation", {}) or {}
        add("הצמדה", "index_uplift", "הצמדת עלויות למדד תשומות הבנייה",
            num(idx.get("build_index_uplift_pct")), "%",
            idx.get("index_base") or idx.get("source", ""))

        for i, area in enumerate(p.get("areas", []) or []):
            label = area.get("label") or area.get("use") or "שטח %d" % (i + 1)
            add("שטחים", "area_sqm_%d" % i, "שטח — %s" % label, num(area.get("sqm")), "מ\"ר",
                area.get("source", ""), SQM)
            add("עלות בנייה", "area_rate_%d" % i, "עלות למ\"ר — %s" % label,
                num(area.get("cost_per_sqm")), "₪/מ\"ר", area.get("source", "") or "cost-benchmarks.md", NIS0)

        parking = p.get("parking", {}) or {}
        add("שטחים", "park_spaces", "מקומות חניה תת-קרקעיים", num(parking.get("spaces")), "יח'",
            parking.get("source", ""), QTY)
        add("עלות בנייה", "park_rate", "עלות למקום חניה", num(parking.get("cost_per_space")), "₪",
            parking.get("source", "") or "cost-benchmarks.md", NIS0)

        for i, extra in enumerate(p.get("direct_other", []) or []):
            add("עלות בנייה", "direct_other_%d" % i, extra.get("label", "עלות ישירה נוספת"),
                num(extra.get("amount")), "₪", extra.get("source", ""), NIS0)

        for i, item in enumerate(p.get("indirect", []) or []):
            label = item.get("label", "עלות עקיפה %d" % (i + 1))
            if item.get("amount") not in (None, ""):
                add("עלויות עקיפות", "ind_%d" % i, label, num(item.get("amount")), "₪",
                    item.get("source", ""), NIS0)
            else:
                add("עלויות עקיפות", "ind_%d" % i, label, num(item.get("pct")), "%",
                    item.get("source", "") or "cost-benchmarks.md")

        # עלויות הדיירים נכנסות כשורות הנחה משלהן, כדי שיהיו ניתנות לשינוי
        # בתא אחד כמו כל פרמטר אחר — ולא רק כסכום מוגמר בגיליון העקיפות.
        for i, line in enumerate(self.m.get("tenants", {}).get("lines", []) or []):
            add("דיירים", "tenant_%d" % i, line["label"], line["amount"], "₪",
                line.get("source", "") or "תקן 21 ס' 4.14(ד)", NIS0)

        tax = p.get("tax", {}) or {}
        for key, label in [("purchase_tax_on_tenant_rights", "מס רכישה בגין רכישת זכויות הדיירים"),
                           ("vat_on_tenant_construction", "מע\"מ על שירותי בנייה לדיירים (שאינו בר-קיזוז)")]:
            if num(tax.get(key)):
                add("דיירים", key, label, num(tax.get(key)), "₪",
                    tax.get(key + "_source", "") or "תקן 21 ס' 4.14(ה)", NIS0)

        add("מיסוי", "vat", "שיעור מע\"מ", num(tax.get("vat_pct"), 18), "%", "tax-rules.md")
        add("מיסוי", "profit_tax", "שיעור מס על הרווח (%s)" %
            ("חברה" if self.m["developer_type"] == "company" else "יחיד"),
            self.m["profit_tax_rate"] * 100, "%", "tax-rules.md")
        add("מיסוי", "betterment", "היטל השבחה", num(tax.get("betterment_levy")), "₪",
            tax.get("betterment_source", ""), NIS0)
        add("מיסוי", "rmi_fee", "דמי היתר רמ\"י", num(tax.get("rmi_permit_fee")), "₪",
            tax.get("rmi_source", ""), NIS0)
        add("בצ\"מ", "cont_pct", "בלתי צפוי מראש (בצ\"מ)", num(p.get("contingency_pct"), 5), "%",
            "cost-benchmarks.md — טווח מקובל 5%-10%")

        rev = p.get("revenue", {}) or {}
        for i, item in enumerate(rev.get("items", []) or []):
            label = item.get("label", "רכיב הכנסה %d" % (i + 1))
            add("הכנסות", "rev_units_%d" % i, "כמות — %s" % label, num(item.get("units"), 1), "יח'",
                item.get("source", ""), QTY)
            add("הכנסות", "rev_sqm_%d" % i, "שטח ממוצע — %s" % label, num(item.get("avg_sqm")), "מ\"ר",
                item.get("source", ""), SQM)
            add("הכנסות", "rev_price_%d" % i, "מחיר מכירה למ\"ר — %s" % label,
                num(item.get("price_per_sqm")), "₪/מ\"ר",
                item.get("source", "") or "עסקאות השוואה — data-sources.md", NIS0)
            if item.get("total_price") not in (None, ""):
                add("הכנסות", "rev_total_%d" % i, "סכום כולל — %s" % label,
                    num(item.get("total_price")), "₪", item.get("source", ""), NIS0)

        ia = p.get("income_asset") or {}
        if ia:
            add("נכס מניב", "noi", "NOI שנתי", num(ia.get("annual_noi")), "₪", ia.get("source", ""), NIS0)
            add("נכס מניב", "cap_rate", "שיעור היוון (Cap Rate)", num(ia.get("cap_rate")), "%",
                ia.get("source", ""))
            add("נכס מניב", "holding_years", "שנות החזקה עד המימוש",
                num(ia.get("holding_years")), "שנים", ia.get("source", ""), QTY)

        fin = p.get("finance", {}) or {}
        add("מימון", "equity", "הון עצמי", num(fin.get("equity")), "₪", "", NIS0)
        add("מימון", "rate", "ריבית אשראי ליווי שנתית", self.m["finance"]["annual_rate"] * 100, "%",
            "ריבית בנק ישראל + מרווח — data-sources.md")
        add("מימון", "fees_pct", "עמלות פתיחה וניהול אשראי", num(fin.get("fees_pct")), "%")
        add("מימון", "warranty_pct", "ערבויות חוק מכר", num(fin.get("warranty_pct")), "%",
            "חוק המכר (דירות) (הבטחת השקעות)")

        add("ספי בקרה", "threshold", "סף רווח יזמי מזערי",
            self.m["results"]["profit_threshold"] * 100, "%",
            "cost-benchmarks.md — דרישת בנק מלווה"
            if self.m["results"]["threshold_basis"] == "margin"
            else "לא חל על נכס מניב — המבחן הוא המרווח")
        if self.m["results"]["threshold_basis"] == "spread":
            add("ספי בקרה", "min_spread", "מרווח תשואה מזערי (נק' בסיס)",
                self.m["results"].get("min_spread_bps") or 150, "נק' בסיס",
                "נורמה בייזום מניב: 150–200", QTY)

        sched = p.get("schedule", {}) or {}
        add("לוח זמנים", "quarters", "משך הפרויקט", num(sched.get("quarters"), 12), "רבעונים", "", QTY)
        return ws

    # -- 2. קרקע ורכישה ----------------------------------------------------

    def build_land(self):
        ws = self.sheet("קרקע ורכישה", [46, 22, 18, 44])
        r = self.title(ws, 1, "קרקע ורכישה", 4)
        r = self.header(ws, r, ["סעיף", "סכום", "בסיס החישוב", "מקור"])
        R = self.ref
        rows = [
            ("מחיר רכישת הקרקע", "=%s" % R["land_price"], "הנחות יסוד"),
            ("מס רכישה", "=%s*%s/100" % (R["land_price"], R["purchase_tax_rate"]), "מחיר הקרקע × שיעור מס רכישה"),
            ("שכ\"ט עו\"ד ועלויות עסקה", "=%s*%s/100" % (R["land_price"], R["legal_pct"]), "מחיר הקרקע × %"),
            ("דמי תיווך", "=%s*%s/100" % (R["land_price"], R["brokerage_pct"]), "מחיר הקרקע × %"),
            ("היטלי פיתוח בגין הקרקע", "=%s" % R["dev_levies"], "הנחות יסוד"),
            ("עלויות קרקע אחרות", "=%s" % R["land_other"], "הנחות יסוד"),
        ]
        first = r
        for label, formula, basis in rows:
            r = self.line(ws, r, [label, formula, basis, ""], fmts=[None, NIS0, None, None])
        r = self.line(ws, r, ["סה\"כ עלות קרקע ורכישה", self._sum("B", first, r - 1), "", ""],
                      fmts=[None, NIS0, None, None], bold=True, fill=TOT_FILL)
        self.ref["TOTAL_LAND"] = "'קרקע ורכישה'!$B$%d" % (r - 1)
        self.note(ws, r + 1,
                  "מס הרכישה על קרקע נגזר משיעור בגיליון ההנחות. סיווג העסקה (קרקע / מסחרי / "
                  "מלאי עסקי) משפיע על השיעור — ראה references/tax-rules.md.", 4)
        return ws

    # -- 3. עלויות בנייה ישירות -------------------------------------------

    def build_direct(self):
        ws = self.sheet("עלויות בנייה ישירות", [40, 16, 18, 20, 34])
        r = self.title(ws, 1, "עלויות בנייה ישירות", 5)
        r = self.note(ws, r, "כל שורה מוצמדת למדד תשומות הבנייה לפי השיעור שבגיליון ההנחות.", 5)
        r = self.header(ws, r, ["רכיב", "שטח / כמות", "עלות ליחידה", "סה\"כ", "מקור"])
        R = self.ref
        first = r
        for i, area in enumerate(self.p.get("areas", []) or []):
            label = area.get("label") or area.get("use") or "שטח %d" % (i + 1)
            r = self.line(ws, r, [
                label,
                "=%s" % R["area_sqm_%d" % i],
                "=%s" % R["area_rate_%d" % i],
                "=B%d*C%d*(1+%s/100)" % (r, r, R["index_uplift"]),
                area.get("source", "") or "cost-benchmarks.md",
            ], fmts=[None, SQM, NIS0, NIS0, None])

        if num((self.p.get("parking") or {}).get("spaces")):
            r = self.line(ws, r, [
                "חניון תת-קרקעי", "=%s" % R["park_spaces"], "=%s" % R["park_rate"],
                "=B%d*C%d*(1+%s/100)" % (r, r, R["index_uplift"]),
                (self.p.get("parking") or {}).get("source", "") or "cost-benchmarks.md",
            ], fmts=[None, QTY, NIS0, NIS0, None])

        for i, extra in enumerate(self.p.get("direct_other", []) or []):
            r = self.line(ws, r, [
                extra.get("label", "עלות ישירה נוספת"), "", "",
                "=%s*(1+%s/100)" % (R["direct_other_%d" % i], R["index_uplift"]),
                extra.get("source", ""),
            ], fmts=[None, None, None, NIS0, None])

        r = self.line(ws, r, ["סה\"כ עלויות בנייה ישירות", "", "", self._sum("D", first, r - 1), ""],
                      fmts=[None, None, None, NIS0, None], bold=True, fill=TOT_FILL)
        self.ref["TOTAL_DIRECT"] = "'עלויות בנייה ישירות'!$D$%d" % (r - 1)
        return ws

    # -- 4. עלויות עקיפות --------------------------------------------------

    def build_indirect(self):
        ws = self.sheet("עלויות עקיפות", [40, 26, 14, 18, 34])
        r = self.title(ws, 1, "עלויות עקיפות", 5)
        r = self.header(ws, r, ["סעיף", "בסיס החישוב", "שיעור", "סכום", "מקור"])
        R = self.ref
        base_ref = {
            "direct": (R["TOTAL_DIRECT"], "עלויות בנייה ישירות"),
            "land": (R["TOTAL_LAND"], "עלות קרקע"),
            "revenue": (R["TOTAL_REVENUE_NET"], "הכנסות נטו"),
            "revenue_gross": (R["TOTAL_REVENUE_GROSS"], "הכנסות ברוטו"),
        }
        first = r
        for i, item in enumerate(self.p.get("indirect", []) or []):
            label = item.get("label", "עלות עקיפה")
            if item.get("amount") not in (None, ""):
                r = self.line(ws, r, [label, "סכום קבוע", "", "=%s" % R["ind_%d" % i],
                                      item.get("source", "")],
                              fmts=[None, None, None, NIS0, None])
            else:
                basis = (item.get("basis") or "direct").strip()
                ref, desc = base_ref.get(basis, base_ref["direct"])
                r = self.line(ws, r, [
                    label, desc, "=%s/100" % R["ind_%d" % i],
                    "=%s*%s/100" % (ref, R["ind_%d" % i]),
                    item.get("source", "") or "cost-benchmarks.md",
                ], fmts=[None, None, PCT, NIS0, None])

        tenants = self.m.get("tenants", {}) or {}
        tax_keys = [k for k in ("purchase_tax_on_tenant_rights", "vat_on_tenant_construction")
                    if k in R]
        if tenants.get("lines") or tax_keys:
            t_first = r
            for i, line in enumerate(tenants.get("lines") or []):
                r = self.line(ws, r, [line["label"], "עלויות טיפול בדיירים", "",
                                      "=%s" % R["tenant_%d" % i],
                                      line.get("source", "") or "תקן 21 ס' 4.14(ד)"],
                              fmts=[None, None, None, NIS0, None])
            for key in tax_keys:
                label = ("מס רכישה בגין רכישת זכויות הדיירים"
                         if key == "purchase_tax_on_tenant_rights"
                         else "מע\"מ על שירותי בנייה לדיירים (שאינו בר-קיזוז)")
                r = self.line(ws, r, [label, "מיסוי פינוי-בינוי", "", "=%s" % R[key],
                                      "תקן 21 ס' 4.14(ה)"],
                              fmts=[None, None, None, NIS0, None])
            r = self.line(ws, r, ["מזה — סה\"כ עלויות הטיפול בדיירים", "", "",
                                  self._sum("D", t_first, r - 1), ""],
                          fmts=[None, None, None, NIS0, None], bold=True, fill=SUB_FILL)
            # שורת הסיכום הזו היא תת-סכום להצגה בלבד; היא לא נכללת בסכום הכולל
            # למטה, אחרת עלויות הדיירים ייספרו פעמיים.
            t_subtotal_row = r - 1
        else:
            t_subtotal_row = None

        if num((self.p.get("tax") or {}).get("betterment_levy")):
            r = self.line(ws, r, ["היטל השבחה", "אומדן תכנוני", "", "=%s" % R["betterment"],
                                  (self.p.get("tax") or {}).get("betterment_source", "")],
                          fmts=[None, None, None, NIS0, None])
        if num((self.p.get("tax") or {}).get("rmi_permit_fee")):
            r = self.line(ws, r, ["דמי היתר רמ\"י", "אומדן", "", "=%s" % R["rmi_fee"],
                                  (self.p.get("tax") or {}).get("rmi_source", "")],
                          fmts=[None, None, None, NIS0, None])

        total_formula = ("=SUM(D%d:D%d)-D%d" % (first, r - 1, t_subtotal_row)
                         if t_subtotal_row else self._sum("D", first, r - 1))
        r = self.line(ws, r, ["סה\"כ עלויות עקיפות", "", "", total_formula, ""],
                      fmts=[None, None, None, NIS0, None], bold=True, fill=TOT_FILL)
        self.ref["TOTAL_INDIRECT"] = "'עלויות עקיפות'!$D$%d" % (r - 1)
        return ws

    # -- 5. הכנסות ---------------------------------------------------------

    def build_revenue(self):
        ws = self.sheet("הכנסות", [34, 12, 14, 18, 20, 20, 30])
        r = self.title(ws, 1, "הכנסות צפויות", 7)
        incl = (self.p.get("revenue") or {}).get("prices_include_vat", True)
        r = self.note(ws, r,
                      "מחירי המכירה הוזנו %s מע\"מ. המודל מציג גם ברוטו וגם נטו; הרווח היזמי "
                      "נמדד על ההכנסה נטו, שכן המע\"מ מועבר לרשות המסים ואינו הכנסת היזם."
                      % ("כוללי" if incl else "ללא"), 7)
        r = self.header(ws, r, ["רכיב", "יחידות", "שטח ממוצע", "מחיר למ\"ר",
                                "סה\"כ ברוטו", "סה\"כ נטו ממע\"מ", "מקור"])
        R = self.ref
        divisor = "(1+%s/100)" % R["vat"] if incl else "1"
        first = r
        for i, item in enumerate((self.p.get("revenue") or {}).get("items", []) or []):
            # פריט שהוזן בסכום כולל (total_price) גובר על מכפלת יחידות×שטח×מחיר.
            # המנוע תמך בזה מלכתחילה והאקסל התעלם, כך שהשניים היו מציגים מספרים
            # שונים לאותו פריט — ואת זה verify_excel היה תופס רק אחרי המסירה.
            if item.get("total_price") not in (None, ""):
                gross = "=%s" % R["rev_total_%d" % i]
            else:
                gross = "=B%d*C%d*D%d" % (r, r, r)
            r = self.line(ws, r, [
                item.get("label", "רכיב הכנסה"),
                "=%s" % R["rev_units_%d" % i],
                "=%s" % R["rev_sqm_%d" % i],
                "=%s" % R["rev_price_%d" % i],
                gross,
                "=E%d/%s" % (r, divisor),
                item.get("source", "") or "עסקאות השוואה",
            ], fmts=[None, QTY, SQM, NIS0, NIS0, NIS0, None])

        ia = self.p.get("income_asset") or {}
        if ia and num(ia.get("annual_noi")) and num(ia.get("holding_years")):
            r = self.line(ws, r, [
                "הכנסות שכירות בתקופת ההחזקה", "=%s" % R["holding_years"],
                num(ia.get("sqm")), "",
                "=%s*%s" % (R["noi"], R["holding_years"]),
                "=E%d/%s" % (r, divisor),
                "NOI שנתי × שנות החזקה",
            ], fmts=[None, QTY, SQM, None, NIS0, NIS0, None])
            # השטח כאן אינו שטח נמכר — אחרת "עלות למ"ר מכור" מתנפח פי שנות
            # ההחזקה. מתקנים ידנית לאפס בעמודת השטח לצורך הסיכום.
            ws.cell(row=r - 1, column=3, value=0).number_format = SQM
        if ia and num(ia.get("annual_noi")) and num(ia.get("cap_rate")):
            r = self.line(ws, r, [
                ia.get("label", "מימוש נכס מניב (Exit)"), 1, num(ia.get("sqm")), "",
                "=%s/(%s/100)" % (R["noi"], R["cap_rate"]),
                "=E%d/%s" % (r, divisor),
                "היוון NOI שנתי בשיעור ההיוון שבהנחות",
            ], fmts=[None, QTY, SQM, None, NIS0, NIS0, None])

        last_data = r - 1
        # עמודה C בשורת הסיכום מחזיקה את סך השטח המכיר (יחידות × שטח ממוצע),
        # שממנו נגזרת "עלות למ"ר מכור". זה השטח שנמכר בפועל — בפינוי-בינוי הוא
        # קטן מהשטח הבנוי, כי דירות התמורה נבנות ואינן נמכרות.
        sqm_word = "שטח בנוי" if self.m.get("kind") == "income" else "שטח מכיר"
        sqm_total = ("=SUMPRODUCT(B%d:B%d,C%d:C%d)" % (first, last_data, first, last_data)
                     if last_data >= first else 0)
        r = self.line(ws, r, ["סה\"כ הכנסות (ו%s)" % sqm_word, "",
                              sqm_total,
                              "",
                              self._sum("E", first, last_data),
                              self._sum("F", first, last_data), ""],
                      fmts=[None, None, SQM, None, NIS0, NIS0, None], bold=True, fill=TOT_FILL)
        self.ref["TOTAL_REVENUE_GROSS"] = "'הכנסות'!$E$%d" % (r - 1)
        self.ref["TOTAL_REVENUE_NET"] = "'הכנסות'!$F$%d" % (r - 1)
        self.ref["SELLABLE_SQM"] = "'הכנסות'!$C$%d" % (r - 1)
        return ws

    # -- 6. בצ"מ ומיסוי ----------------------------------------------------

    def build_contingency_tax(self):
        ws = self.sheet("בצמ ומיסוי", [40, 26, 14, 20, 34])
        r = self.title(ws, 1, "בצ\"מ ומיסוי", 5)
        r = self.header(ws, r, ["סעיף", "בסיס", "שיעור", "סכום", "הערה"])
        R = self.ref
        base_key = (self.p.get("contingency_base") or "direct_and_indirect").strip()
        base_formula, base_desc = {
            "direct": (R["TOTAL_DIRECT"], "עלויות ישירות"),
            "direct_and_indirect": ("(%s+%s)" % (R["TOTAL_DIRECT"], R["TOTAL_INDIRECT"]),
                                    "עלויות ישירות + עקיפות"),
            "total": ("(%s+%s+%s)" % (R["TOTAL_LAND"], R["TOTAL_DIRECT"], R["TOTAL_INDIRECT"]),
                      "סך העלויות כולל קרקע"),
        }.get(base_key, ("(%s+%s)" % (R["TOTAL_DIRECT"], R["TOTAL_INDIRECT"]), "ישירות + עקיפות"))

        r = self.line(ws, r, ["בלתי צפוי מראש (בצ\"מ)", base_desc, "=%s/100" % R["cont_pct"],
                              "=%s*%s/100" % (base_formula, R["cont_pct"]),
                              "טווח מקובל 5%-10%; בנק מלווה בוחן סעיף זה"],
                      fmts=[None, None, PCT, NIS0, None])
        self.ref["CONTINGENCY"] = "'בצמ ומיסוי'!$D$%d" % (r - 1)

        r += 1
        r = self.line(ws, r, ["מע\"מ", "", "=%s/100" % R["vat"], "",
                              "מע\"מ תשומות מתקזז אצל עוסק מורשה; מחירי מכירה למגורים כוללים מע\"מ"],
                      fmts=[None, None, PCT, None, None])
        r = self.line(ws, r, ["שיעור מס על הרווח",
                              "חברה" if self.m["developer_type"] == "company" else "יחיד",
                              "=%s/100" % R["profit_tax"], "",
                              "יזם הבונה למכירה ממוסה כהכנסה פירותית (מלאי עסקי), לא כמס שבח — "
                              "ראה references/tax-rules.md"],
                      fmts=[None, None, PCT, None, None])
        return ws

    # -- 7. מימון ----------------------------------------------------------

    def build_finance(self):
        ws = self.sheet("מימון", [40, 26, 14, 20, 34])
        r = self.title(ws, 1, "מימון", 5)
        r = self.note(ws, r,
                      "הריבית מחושבת בגיליון התזרים על היתרה המנוצלת בפועל בכל רבעון, "
                      "לא כמכפלה גסה של קרן בזמן.", 5)
        r = self.header(ws, r, ["סעיף", "בסיס", "שיעור", "סכום", "הערה"])
        R = self.ref
        r = self.line(ws, r, ["הון עצמי", "", "", "=%s" % R["equity"], "נספג ראשון בתזרים"],
                      fmts=[None, None, None, NIS0, None])
        r = self.line(ws, r, ["שיא ניצול אשראי ליווי", "מתוך התזרים", "",
                              "=MAX(%s)" % self.ref["CF_BALANCE_RANGE"], "מסגרת האשראי הנדרשת"],
                      fmts=[None, None, None, NIS0, None])
        self.ref["PEAK_DEBT"] = "'מימון'!$D$%d" % (r - 1)
        first = r
        r = self.line(ws, r, ["ריבית אשראי ליווי", "יתרה מנוצלת רבעונית", "=%s/100" % R["rate"],
                              "=SUM(%s)" % self.ref["CF_INTEREST_RANGE"], "מצטבר על פני חיי הפרויקט"],
                      fmts=[None, None, PCT, NIS0, None])
        r = self.line(ws, r, ["עמלות פתיחה וניהול אשראי", "מסגרת האשראי", "=%s/100" % R["fees_pct"],
                              "=%s*%s/100" % (self.ref["PEAK_DEBT"], R["fees_pct"]), ""],
                      fmts=[None, None, PCT, NIS0, None])
        r = self.line(ws, r, ["ערבויות חוק מכר", "הכנסות ברוטו", "=%s/100" % R["warranty_pct"],
                              "=%s*%s/100" % (R["TOTAL_REVENUE_GROSS"], R["warranty_pct"]),
                              "חוק המכר (דירות) (הבטחת השקעות)"],
                      fmts=[None, None, PCT, NIS0, None])
        r = self.line(ws, r, ["סה\"כ עלויות מימון", "", "", self._sum("D", first, r - 1), ""],
                      fmts=[None, None, None, NIS0, None], bold=True, fill=TOT_FILL)
        self.ref["TOTAL_FINANCE"] = "'מימון'!$D$%d" % (r - 1)
        return ws

    # -- 8. תזרים רבעוני ---------------------------------------------------

    def build_cashflow(self):
        cf = self.m["cashflow"]
        n = len(cf)
        ws = self.sheet("תזרים רבעוני", [14, 13, 16, 16, 16, 16, 16, 16, 16, 16, 17, 17, 19])
        r = self.title(ws, 1, "תזרים מזומנים רבעוני", 12)
        r = self.note(ws, r,
                      "התאריכים הם תאריכי Excel אמיתיים (ניתן למיין ולסנן לפיהם), לא טקסט. "
                      "יתרת האשראי נבנית מרבעון לרבעון: הון עצמי נספג ראשון, והריבית נצברת "
                      "רק על מה שמעבר לו.", 12)
        r = self.header(ws, r, ["רבעון", "תאריך", "קרקע", "בנייה ישירה", "עקיפות", "בצ\"מ",
                                "סה\"כ הוצאות", "הכנסות", "ניצול הון עצמי", "יתרת הון עצמי",
                                "ריבית", "יתרת אשראי", "תזרים הון עצמי (ל-IRR)"])
        R = self.ref
        first = r
        sched = self.p.get("schedule", {}) or {}

        def weights(key, fallback=None):
            w = sched.get(key) or fallback or []
            w = [num(x) for x in w][:n]
            w += [0.0] * (n - len(w))
            s = sum(w)
            return [x / s if s else 1.0 / n for x in w]

        land_w = weights("land_curve", [1] + [0] * (n - 1))
        cost_w = weights("cost_curve", [1] * n)
        ind_w = weights("indirect_curve", sched.get("cost_curve") or [1] * n)
        rev_w = weights("revenue_curve", [1] * n)

        for i, row in enumerate(cf):
            rr = first + i
            prev = rr - 1
            d = date.fromisoformat(row["date"])
            need = "MAX(0,G{r}-H{r})".format(r=rr)
            if i == 0:
                eq_avail = R["equity"]
            else:
                eq_avail = "J%d" % prev
            bal_prev = "0" if i == 0 else "L%d" % prev
            self.line(ws, rr, [
                row["quarter"], d,
                "=%s*%s" % (R["TOTAL_LAND"], _fmt(land_w[i])),
                "=%s*%s" % (R["TOTAL_DIRECT"], _fmt(cost_w[i])),
                "=%s*%s" % (R["TOTAL_INDIRECT"], _fmt(ind_w[i])),
                "=%s*%s" % (R["CONTINGENCY"], _fmt(cost_w[i])),
                "=SUM(C%d:F%d)" % (rr, rr),
                "=%s*%s" % (R["TOTAL_REVENUE_NET"], _fmt(rev_w[i])),
                "=MIN(%s,%s)" % (eq_avail, need),
                "=%s-I%d" % (eq_avail, rr),
                "=MAX(0,%s-(H%d-G%d+I%d))*%s/400" % (bal_prev, rr, rr, rr, R["rate"]),
                "=MAX(0,%s-(H%d-G%d+I%d))+K%d" % (bal_prev, rr, rr, rr, rr),
                # החלוקה ליזם ברבעון האחרון מוזרקת ב-_inject_terminal_flow אחרי
                # שגיליון "תוצאות" נבנה ומיקום תא הרווח-לאחר-מס ידוע בוודאות.
                "=-I%d" % rr,
            ], fmts=[None, DATE_FMT, NIS0, NIS0, NIS0, NIS0, NIS0, NIS0, NIS0, NIS0, NIS0, NIS0, NIS0])

        last = first + n - 1
        self.line(ws, last + 1, ["סה\"כ", "",
                                 "=SUM(C%d:C%d)" % (first, last), "=SUM(D%d:D%d)" % (first, last),
                                 "=SUM(E%d:E%d)" % (first, last), "=SUM(F%d:F%d)" % (first, last),
                                 "=SUM(G%d:G%d)" % (first, last), "=SUM(H%d:H%d)" % (first, last),
                                 "=SUM(I%d:I%d)" % (first, last), "",
                                 "=SUM(K%d:K%d)" % (first, last), "=MAX(L%d:L%d)" % (first, last),
                                 "=SUM(M%d:M%d)" % (first, last)],
                  fmts=[None, None] + [NIS0] * 11, bold=True, fill=TOT_FILL)
        self.ref["CF_INTEREST_RANGE"] = "'תזרים רבעוני'!$K$%d:$K$%d" % (first, last)
        self.ref["CF_BALANCE_RANGE"] = "'תזרים רבעוני'!$L$%d:$L$%d" % (first, last)
        self.ref["CF_EQUITY_RANGE"] = "'תזרים רבעוני'!$I$%d:$I$%d" % (first, last)
        self.ref["CF_IRR_RANGE"] = "'תזרים רבעוני'!$M$%d:$M$%d" % (first, last)
        self._terminal = (ws, last, first)
        return ws

    # -- 9. תוצאות ---------------------------------------------------------

    def build_results(self):
        ws = self.sheet("תוצאות", [42, 24, 52])
        r = self.title(ws, 1, "תוצאות המודל", 3)
        r = self.header(ws, r, ["מדד", "ערך", "הסבר"])
        R = self.ref
        res = self.m["results"]

        cost_formula = "(%s+%s+%s+%s+%s)" % (R["TOTAL_LAND"], R["TOTAL_DIRECT"],
                                             R["TOTAL_INDIRECT"], R["CONTINGENCY"],
                                             R["TOTAL_FINANCE"])
        r = self.line(ws, r, ["סך עלויות הפרויקט", "=" + cost_formula,
                              "קרקע + בנייה ישירה + עקיפות + בצ\"מ + מימון"],
                      fmts=[None, NIS0, None], bold=True)
        self.ref["TOTAL_COST"] = "'תוצאות'!$B$%d" % (r - 1)
        r = self.line(ws, r, ["סך הכנסות (ברוטו)", "=%s" % R["TOTAL_REVENUE_GROSS"], "כולל מע\"מ"],
                      fmts=[None, NIS0, None])
        r = self.line(ws, r, ["סך הכנסות (נטו ממע\"מ)", "=%s" % R["TOTAL_REVENUE_NET"],
                              "ההכנסה שנותרת בידי היזם"], fmts=[None, NIS0, None], bold=True)
        r += 1
        r = self.line(ws, r, ["רווח יזמי לפני מס",
                              "=%s-%s" % (R["TOTAL_REVENUE_NET"], self.ref["TOTAL_COST"]),
                              "הכנסות נטו פחות סך העלויות"],
                      fmts=[None, NIS0, None], bold=True)
        profit_ref = "'תוצאות'!$B$%d" % (r - 1)
        self.ref["PROFIT"] = profit_ref
        r = self.line(ws, r, ["רווח יזמי — % מהעלויות",
                              "=IFERROR(%s/%s,0)" % (profit_ref, self.ref["TOTAL_COST"]),
                              "המדד שבנק מלווה בוחן. סף מקובל: 15%-18%"],
                      fmts=[None, PCT, None], bold=True)
        margin_ref = "'תוצאות'!$B$%d" % (r - 1)
        self.ref["MARGIN"] = margin_ref
        r = self.line(ws, r, ["רווח יזמי — % מההכנסות",
                              "=IFERROR(%s/%s,0)" % (profit_ref, R["TOTAL_REVENUE_NET"]),
                              "מדד משלים; נמוך מהמדד על העלויות מטבעו"],
                      fmts=[None, PCT, None])
        r = self.line(ws, r, ["מס על הרווח",
                              "=MAX(0,%s)*%s/100" % (profit_ref, R["profit_tax"]),
                              "לפי סיווג היזם"], fmts=[None, NIS0, None])
        tax_ref = "'תוצאות'!$B$%d" % (r - 1)
        r = self.line(ws, r, ["רווח לאחר מס", "=%s-%s" % (profit_ref, tax_ref), ""],
                      fmts=[None, NIS0, None], bold=True)
        after_tax = "'תוצאות'!$B$%d" % (r - 1)
        self.ref["PROFIT_AFTER_TAX"] = after_tax
        r += 1
        # IRR חי מתזרים ההון העצמי, מוצג שנתי
        n = len(self.m["cashflow"])
        r = self.line(ws, r, ["IRR שנתי על ההון העצמי",
                              "=IFERROR((1+IRR(%s))^4-1,\"לא ניתן לחישוב\")" % self.ref["CF_IRR_RANGE"],
                              "מחושב מתזרים ההון העצמי הרבעוני ומוצג במונחים שנתיים"],
                      fmts=[None, PCT, None], bold=True)
        r = self.line(ws, r, ["תשואה על ההון (ROC)",
                              "=IFERROR(%s/%s,\"\")" % (after_tax, R["equity"]),
                              "רווח לאחר מס חלקי ההון העצמי"], fmts=[None, PCT, None])
        sqm_label = "עלות למ\"ר בנוי" if self.m.get("kind") == "income" else "עלות למ\"ר מכור"
        r = self.line(ws, r, [sqm_label,
                              "=IFERROR(%s/%s,\"\")" % (self.ref["TOTAL_COST"], R["SELLABLE_SQM"]),
                              "סך העלויות חלקי השטח המכיר"], fmts=[None, NIS0, None])
        r = self.line(ws, r, ["שיא ניצול אשראי", "=%s" % self.ref["PEAK_DEBT"],
                              "מסגרת הליווי הנדרשת"], fmts=[None, NIS0, None])
        r += 1
        im = self.m.get("income_metrics")
        if self.m["results"]["threshold_basis"] == "spread" and im:
            # בנכס מניב המבחן הוא המרווח. הצגת "סף רווח יזמי" כאן סותרת את
            # הדוח ומטעה את מי שקורא רק את האקסל.
            r = self.line(ws, r, ["תשואה על העלות (Yield on Cost)",
                                  "=%s/%s" % (R["noi"], self.ref["TOTAL_COST"]),
                                  "NOI שנתי חלקי סך עלות הפרויקט"],
                          fmts=[None, '0.00%', None], bold=True)
            yoc_ref = "'תוצאות'!$B$%d" % (r - 1)
            r = self.line(ws, r, ["שיעור היוון ביציאה (Exit Cap)", "=%s/100" % R["cap_rate"],
                                  "מגיליון ההנחות"], fmts=[None, '0.00%', None])
            r = self.line(ws, r, ["מרווח (נק' בסיס)",
                                  "=(%s-%s/100)*10000" % (yoc_ref, R["cap_rate"]),
                                  "נורמה בייזום מניב: 150–200"], fmts=[None, '#,##0', None], bold=True)
            spread_ref = "'תוצאות'!$B$%d" % (r - 1)
            r = self.line(ws, r, ["שיעור היוון לאיזון", "=%s" % yoc_ref,
                                  "מעליו הפרויקט מפסיד"], fmts=[None, '0.00%', None])
            r = self.line(ws, r, ["מרווח מזערי נדרש", "=%s" % R["min_spread"],
                                  "מגיליון ההנחות"], fmts=[None, '#,##0', None])
            r = self.line(ws, r, ["עמידה במבחן המרווח",
                                  "=IF(%s>=%s,\"עומד במבחן\",\"אינו עומד במבחן\")"
                                  % (spread_ref, R["min_spread"]),
                                  "בנכס מניב זהו המבחן הקובע, לא סף הרווח היזמי"],
                          fmts=[None, None, None], bold=True)
            check_cell = ws.cell(row=r - 1, column=2)
            ws.conditional_formatting.add(check_cell.coordinate, CellIsRule(
                operator="equal", formula=['"עומד במבחן"'],
                fill=PatternFill("solid", fgColor=GREEN)))
            ws.conditional_formatting.add(check_cell.coordinate, CellIsRule(
                operator="equal", formula=['"אינו עומד במבחן"'],
                fill=PatternFill("solid", fgColor=RED)))
            self.note(ws, r + 1, DISCLAIMER, 3)
            return ws

        r = self.line(ws, r, ["סף רווח יזמי נדרש", "=%s/100" % R["threshold"], "מגיליון ההנחות"],
                      fmts=[None, PCT, None])
        r = self.line(ws, r, ["עמידה בסף",
                              "=IF(%s>=%s/100,\"עומד בסף\",\"מתחת לסף\")" % (margin_ref, R["threshold"]),
                              "מתעדכן אוטומטית עם שינוי ההנחות"], fmts=[None, None, None], bold=True)
        check_cell = ws.cell(row=r - 1, column=2)
        ws.conditional_formatting.add(
            check_cell.coordinate,
            CellIsRule(operator="equal", formula=['"עומד בסף"'],
                       fill=PatternFill("solid", fgColor=GREEN)))
        ws.conditional_formatting.add(
            check_cell.coordinate,
            CellIsRule(operator="equal", formula=['"מתחת לסף"'],
                       fill=PatternFill("solid", fgColor=RED)))
        self.note(ws, r + 1, DISCLAIMER, 3)
        return ws

    # -- 10. ניתוחי רגישות -------------------------------------------------

    def build_sensitivity(self):
        ws = self.sheet("ניתוחי רגישות", [26] + [15] * 8)
        r = self.title(ws, 1, "ניתוחי רגישות", 8)
        r = self.note(ws, r,
                      "הערכים בגיליון זה מחושבים במנוע המודל (כל תא = הרצה מלאה של המודל כולל "
                      "לולאת המימון) ולכן אינם נוסחאות חיות. אחרי שינוי הנחות — הרץ מחדש את "
                      "build_excel.py כדי לרענן את הטבלאות האלה. שאר הגיליונות מתעדכנים לבד.", 8)

        sens = sensitivity_grid(self.p)
        th = self.m["results"]["profit_threshold"]

        r = self.line(ws, r, ["טבלה 1 — רווח יזמי (% מהעלויות): מחיר מכירה × עלות בנייה"],
                      bold=True)
        hdr = ["עלות בנייה \\ מחיר מכירה"] + ["%+d%%" % s for s in sens["price_steps"]]
        r = self.header(ws, r, hdr)
        top = r
        for i, cs in enumerate(sens["cost_steps"]):
            vals = ["%+d%%" % cs] + sens["grid"][i]
            r = self.line(ws, r, vals,
                          fmts=[None] + [PCT] * len(sens["price_steps"]))
        rng = "B%d:%s%d" % (top, get_column_letter(1 + len(sens["price_steps"])), r - 1)
        # ירוק מעל הסף, צהוב בטווח 3 נקודות מתחתיו, אדום מתחת לכך.
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="greaterThanOrEqual", formula=[str(th)],
            fill=PatternFill("solid", fgColor=GREEN)))
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="between", formula=[str(max(0.0, th - 0.03)), str(th)],
            fill=PatternFill("solid", fgColor=YELLOW)))
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="lessThan", formula=[str(max(0.0, th - 0.03))],
            fill=PatternFill("solid", fgColor=RED)))
        r = self.note(ws, r, "ירוק = עומד בסף %.0f%%; צהוב = עד 3 נקודות מתחת לסף; אדום = מתחת לכך."
                      % (th * 100), 8)

        sensp = sensitivity_grid(self.p, metric="profit_before_tax")
        r = self.line(ws, r, ["טבלה 2 — רווח יזמי (₪): מחיר מכירה × עלות בנייה"], bold=True)
        r = self.header(ws, r, hdr)
        top2 = r
        for i, cs in enumerate(sensp["cost_steps"]):
            r = self.line(ws, r, ["%+d%%" % cs] + sensp["grid"][i],
                          fmts=[None] + [NIS0] * len(sensp["price_steps"]))
        rng2 = "B%d:%s%d" % (top2, get_column_letter(1 + len(sensp["price_steps"])), r - 1)
        ws.conditional_formatting.add(rng2, CellIsRule(
            operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=RED)))
        r += 1

        r = self.line(ws, r, ["טבלה 3 — תרחישי ריבית"], bold=True)
        r = self.header(ws, r, ["שינוי בריבית", "ריבית שנתית", "עלות ריבית",
                                "רווח יזמי (₪)", "רווח יזמי (% מהעלויות)"])
        top3 = r
        for sc in rate_scenarios(self.p):
            r = self.line(ws, r, ["%+.1f נק' אחוז" % sc["delta_pp"], sc["rate"],
                                  sc["interest"], sc["profit"], sc["metric"]],
                          fmts=[None, PCT, NIS0, NIS0, PCT])
        ws.conditional_formatting.add("E%d:E%d" % (top3, r - 1), CellIsRule(
            operator="lessThan", formula=[str(th)], fill=PatternFill("solid", fgColor=RED)))
        return ws

    # -- 12. תמורות לדיירים (תקן 21 חלקים ב'-ג') ----------------------------

    def build_tenant_consideration(self):
        """גיליון התמורות. קיים רק כשהוזנו נתוני דירות — הוא לא ממציא שוויים.

        השוויים כאן הם קלט של שמאי, לא תוצר של המודל, ולכן הם נכתבים כערכים
        ולא כנוסחאות. מה שכן מחושב הוא הפער והיחס, והם נוסחאות חיות כדי שאפשר
        יהיה לשחק בשווי ולראות את היחס זז."""
        from model import tenant_consideration
        tc = tenant_consideration(self.p)
        if not tc["typical"] and not tc["specific"]:
            return None
        ws = self.sheet("תמורות לדיירים", [30, 10, 14, 20, 14, 20, 18, 20, 10, 26])
        r = self.title(ws, 1, "תמורות לדיירים — תקן 21 חלקים ב' ו-ג'", 10)
        r = self.note(ws, r,
                      "השוויים לפני ואחרי הם קלט של שמאי מקרקעין ואינם נקבעים על ידי "
                      "הכלי. הפער והיחס מחושבים בנוסחה חיה.", 10)

        for title, rows, with_extras in [
                ("חלק ב' — דירות אופייניות (ס' 5.11–5.13)", tc["typical"], False),
                ("חלק ג' — דירות מסוימות (ס' 6.11–6.13)", tc["specific"], True)]:
            if not rows:
                continue
            r = self.line(ws, r, [title], bold=True)
            r = self.header(ws, r, ["דירה", "כמות", "שטח קיים", "שווי לפני", "שטח חדש",
                                    "שווי אחרי", "הצמדות", "פער", "יחס", "מקור"])
            first = r
            for t in rows:
                r = self.line(ws, r, [
                    t["label"], t["count"], t["existing_sqm"], t["existing_value"],
                    t["new_sqm"], t["new_value"], t["extras_total"],
                    "=F%d+G%d-D%d" % (r, r, r),
                    "=IFERROR((F%d+G%d)/D%d,\"\")" % (r, r, r),
                    t.get("source") or "—",
                ], fmts=[None, QTY, SQM, NIS0, SQM, NIS0, NIS0, NIS0, '0.00', None])
            r = self.line(ws, r, ["סה\"כ", "=SUM(B%d:B%d)" % (first, r - 1), "",
                                  "=SUMPRODUCT(B%d:B%d,D%d:D%d)" % (first, r - 1, first, r - 1),
                                  "", "=SUMPRODUCT(B%d:B%d,F%d:F%d)" % (first, r - 1, first, r - 1),
                                  "", "=SUMPRODUCT(B%d:B%d,H%d:H%d)" % (first, r - 1, first, r - 1),
                                  "", ""],
                          fmts=[None, QTY, None, NIS0, None, NIS0, None, NIS0, None, None],
                          bold=True, fill=TOT_FILL)
            r += 1
        self.note(ws, r, DISCLAIMER, 10)
        return ws

    # -- 11. דגלים אדומים --------------------------------------------------

    def build_flags(self):
        ws = self.sheet("דגלים אדומים", [14, 34, 88])
        r = self.title(ws, 1, "דגלים אדומים", 3)
        r = self.note(ws, r,
                      "הדגלים נגזרים מהמודל ולא נכתבו ביד — הרצה מחדש עם מספרים חדשים מפיקה "
                      "רשימה חדשה. דגל אינו פסילה של הפרויקט; הוא נקודה שדורשת החלטה מודעת.", 3)
        r = self.header(ws, r, ["חומרה", "נושא", "פירוט"])
        flags = derive_flags(self.p, self.m)
        if not flags:
            r = self.line(ws, r, ["—", "לא אותרו דגלים",
                                  "המודל לא זיהה חריגה מהספים שהוגדרו. אין בכך אישור לפרויקט."])
        for f in flags:
            level = "אדום" if f.get("level") == "red" else "צהוב"
            r = self.line(ws, r, [level, f.get("title", ""), f.get("text", "")],
                          fill=PatternFill("solid", fgColor=RED if level == "אדום" else YELLOW))

        devs = deviation_notes(self.p)
        if devs:
            r += 1
            r = self.line(ws, r, ["סטיות מאסמכתאות חיצוניות (מעל 10%)"], bold=True)
            r = self.header(ws, r, ["חומרה", "פרמטר", "פירוט"])
            for d in devs:
                r = self.line(ws, r, ["סטייה", d["label"], d["text"]],
                              fill=PatternFill("solid", fgColor=YELLOW))
        self.note(ws, r + 1, DISCLAIMER, 3)
        return ws

    # -- הרכבה -------------------------------------------------------------

    def build(self, path):
        # סדר הבנייה נקבע לפי תלויות הנוסחאות, לא לפי סדר התצוגה: גיליון שמפנה
        # לטווח חייב שהטווח כבר יירשם ב-self.ref. בסוף מסדרים מחדש לסדר הקריאה.
        self.build_assumptions()
        self.build_land()
        self.build_direct()
        self.build_revenue()          # לפני העקיפות — הן מפנות להכנסות
        self.build_indirect()
        self.build_contingency_tax()
        self._pre_register_cashflow_ranges()
        self.build_cashflow()
        self.build_finance()
        self.build_results()
        self._inject_terminal_flow()
        self.build_sensitivity()
        self.build_tenant_consideration()
        self.build_flags()
        self._reorder()
        self.wb.save(path)
        return path

    def _pre_register_cashflow_ranges(self):
        """התזרים מפנה ל-CONTINGENCY ו-TOTAL_* שכבר קיימים, אבל 'מימון' ו'תוצאות'
        מפנים לטווחי התזרים. כאן נרשמות הכתובות מראש כדי לשבור את מעגל התלות —
        המיקומים ידועים מראש כי מספר הרבעונים ידוע."""
        n = len(self.m["cashflow"])
        first = 6   # 1 כותרת, 2 ריק, 3 הערה, 4 ריק, 5 כותרות טבלה → נתונים מ-6
        last = first + n - 1
        self.ref["CF_INTEREST_RANGE"] = "'תזרים רבעוני'!$K$%d:$K$%d" % (first, last)
        self.ref["CF_BALANCE_RANGE"] = "'תזרים רבעוני'!$L$%d:$L$%d" % (first, last)
        self.ref["CF_IRR_RANGE"] = "'תזרים רבעוני'!$M$%d:$M$%d" % (first, last)

    def _inject_terminal_flow(self):
        """מזריק את החלוקה ליזם לתא האחרון בעמודת תזרים ההון העצמי.

        זה נעשה במעבר שני ולא בזמן בניית התזרים, מפני שהחלוקה מפנה לתא
        "רווח לאחר מס" שנמצא בגיליון התוצאות — גיליון שנבנה אחרי התזרים.
        ניסיון קודם לנחש את מספר השורה מראש פספס בשורה אחת, וה-IRR באקסל יצא
        אפס בזמן שהמודל בפייתון הראה 19.8%. הזרקה בדיעבד מסירה את הניחוש:
        הכתובת נלקחת מ-self.ref שנרשם בפועל.
        """
        ws, last, first = self._terminal
        ws.cell(row=last, column=13,
                value="=-I%d+SUM(I%d:I%d)+%s" % (last, first, last, self.ref["PROFIT_AFTER_TAX"]))

    def _reorder(self):
        order = [ASSUMPTIONS_SHEET, "קרקע ורכישה", "עלויות בנייה ישירות", "עלויות עקיפות",
                 "מימון", "בצמ ומיסוי", "הכנסות", "תזרים רבעוני", "תוצאות",
                 "ניתוחי רגישות", "תמורות לדיירים", "דגלים אדומים"]
        self.wb._sheets = [self.wb[name] for name in order if name in self.wb.sheetnames]


def _fmt(x):
    return ("%.10f" % x).rstrip("0").rstrip(".") or "0"


def main():
    ap = argparse.ArgumentParser(description="בונה את המודל הכלכלי כקובץ Excel")
    ap.add_argument("project")
    ap.add_argument("-o", "--output", default="מודל כלכלי.xlsx")
    args = ap.parse_args()
    project = load_project(args.project)
    path = Builder(project).build(args.output)
    print("נוצר: %s" % path)
    return 0


if __name__ == "__main__":
    sys.exit(main())

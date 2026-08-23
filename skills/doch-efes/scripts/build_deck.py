#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_deck.py — בונה מצגת PPTX ליזם, RTL.

המצגת אינה תקציר של הדוח אלא כלי החלטה: 8–12 שקפים שמובילים מהפרויקט אל
ההמלצה. השקף הקובע הוא שקף הרווחיות; הרגישויות והדגלים קיימים כדי שהיזם יראה
מה מפיל את התשובה, לא כדי למלא מקום.

RTL ב-python-pptx: אין דגל ברמת המצגת. מסמנים כל פסקה ב-``a:pPr rtl="1"``
ומיישרים לימין. מספרים ומונחים לועזיים בתוך שורה עברית מסתדרים לבד — אין
להפוך תווים ידנית.

שימוש:
    python3 build_deck.py project.json -o "מצגת ליזם.pptx"
"""

from __future__ import annotations

import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from model import DISCLAIMER, full_output, load_project  # noqa: E402

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Inches, Pt
except ImportError:  # pragma: no cover
    print("חסרה תלות. התקן:  pip install python-pptx", file=sys.stderr)
    raise

FONT = "Arial"
NAVY = RGBColor(0x1F, 0x38, 0x64)
ACCENT = RGBColor(0x2E, 0x5C, 0x8A)
GREEN = RGBColor(0x1E, 0x7A, 0x46)
RED = RGBColor(0xC0, 0x00, 0x00)
AMBER = RGBColor(0xB8, 0x86, 0x0B)
GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xEE, 0xF3, 0xFA)

W, H = Inches(13.333), Inches(7.5)


def money(x, short=False):
    try:
        v = float(x)
    except (TypeError, ValueError):
        return "—"
    if short and abs(v) >= 1_000_000:
        return "%.1fM ₪" % (v / 1_000_000)
    return "%s ₪" % format(int(round(v)), ",")


def percent(x, digits=1):
    try:
        return ("%." + str(digits) + "f%%") % (float(x) * 100)
    except (TypeError, ValueError):
        return "—"


def rtl(paragraph, align=PP_ALIGN.RIGHT):
    """מסמן פסקה כ-RTL ברמת ה-XML. בלי זה PowerPoint מיישר שמאלה."""
    paragraph.alignment = align
    pPr = paragraph._pPr if paragraph._pPr is not None else paragraph._p.get_or_add_pPr()
    pPr.set("rtl", "1")
    return paragraph


class Deck:
    def __init__(self, data):
        self.d = data
        self.m = data["model"]
        self.kind = data["model"].get("kind", "sale")
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]

    # -- אבני בניין --------------------------------------------------------

    def slide(self, title=None):
        s = self.prs.slides.add_slide(self.blank)
        if title:
            box = s.shapes.add_textbox(Inches(0.6), Inches(0.35), W - Inches(1.2), Inches(0.9))
            tf = box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            rtl(p)
            run = p.add_run()
            run.text = title
            run.font.size = Pt(30)
            run.font.bold = True
            run.font.name = FONT
            run.font.color.rgb = NAVY
            line = s.shapes.add_shape(1, Inches(0.6), Inches(1.28), W - Inches(1.2), Pt(2.5))
            line.fill.solid()
            line.fill.fore_color.rgb = ACCENT
            line.line.fill.background()
            line.shadow.inherit = False
        return s

    def text(self, slide, text, top, size=16, bold=False, color=None,
             left=Inches(0.6), width=None, align=PP_ALIGN.RIGHT):
        box = slide.shapes.add_textbox(left, top, width or (W - Inches(1.2)), Inches(0.5))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        rtl(p, align)
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = FONT
        run.font.color.rgb = color or RGBColor(0x22, 0x22, 0x22)
        return box

    def bullets(self, slide, items, top, size=17):
        box = slide.shapes.add_textbox(Inches(0.8), top, W - Inches(1.6), H - top - Inches(0.8))
        tf = box.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            rtl(p)
            p.space_after = Pt(10)
            run = p.add_run()
            run.text = "•  " + item
            run.font.size = Pt(size)
            run.font.name = FONT
        return box

    def kpi_row(self, slide, cards, top=Inches(2.0), height=Inches(1.7)):
        """שורת כרטיסי מדד. הרוחב מחולק שווה בשווה כדי שהשקף יישאר מאוזן
        בין 3 ל-4 כרטיסים בלי כוונון ידני."""
        n = len(cards)
        margin, gap = Inches(0.6), Inches(0.25)
        total = W - 2 * margin - gap * (n - 1)
        cw = int(total / n)
        for i, (label, value, color) in enumerate(cards):
            # RTL: הכרטיס הראשון מימין
            left = int(W - margin - cw - i * (cw + gap))
            card = slide.shapes.add_shape(5, left, top, cw, height)
            card.fill.solid()
            card.fill.fore_color.rgb = LIGHT
            card.line.color.rgb = RGBColor(0xD0, 0xDC, 0xEC)
            card.shadow.inherit = False
            tf = card.text_frame
            tf.word_wrap = True
            tf.margin_top = Inches(0.18)
            p = tf.paragraphs[0]
            rtl(p, PP_ALIGN.CENTER)
            r = p.add_run()
            r.text = value
            r.font.size = Pt(26)
            r.font.bold = True
            r.font.name = FONT
            r.font.color.rgb = color or NAVY
            p2 = tf.add_paragraph()
            rtl(p2, PP_ALIGN.CENTER)
            r2 = p2.add_run()
            r2.text = label
            r2.font.size = Pt(13)
            r2.font.name = FONT
            r2.font.color.rgb = GREY

    def table(self, slide, headers, rows, top, col_widths=None, size=13):
        n_rows, n_cols = len(rows) + 1, len(headers)
        height = min(Inches(0.42) * n_rows, H - top - Inches(0.5))
        shape = slide.shapes.add_table(n_rows, n_cols, Inches(0.6), top,
                                       W - Inches(1.2), height)
        tbl = shape.table
        # RTL: הופכים את סדר העמודות כדי שהעמודה הראשונה תופיע מימין
        headers = list(reversed(headers))
        rows = [list(reversed(r)) for r in rows]
        if col_widths:
            widths = list(reversed(col_widths))
            total = sum(widths)
            for i, w in enumerate(widths):
                tbl.columns[i].width = Emu(int((W - Inches(1.2)) * w / total))
        for c, h in enumerate(headers):
            cell = tbl.cell(0, c)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            rtl(p, PP_ALIGN.CENTER)
            r = p.add_run()
            r.text = str(h)
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.name = FONT
            r.font.color.rgb = WHITE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY
        for ri, row in enumerate(rows, start=1):
            for ci, val in enumerate(row):
                cell = tbl.cell(ri, ci)
                cell.text = ""
                p = cell.text_frame.paragraphs[0]
                # העמודה האחרונה (אחרי ההיפוך) היא עמודת התוויות → יישור לימין
                rtl(p, PP_ALIGN.RIGHT if ci == n_cols - 1 else PP_ALIGN.CENTER)
                r = p.add_run()
                r.text = str(val)
                r.font.size = Pt(size)
                r.font.name = FONT
        return tbl

    def footer(self, slide):
        self.text(slide, DISCLAIMER, H - Inches(0.55), size=9, color=GREY,
                  align=PP_ALIGN.CENTER)

    # -- השקפים ------------------------------------------------------------

    def s_cover(self):
        s = self.slide()
        meta = self.m["meta"]
        bg = s.shapes.add_shape(1, 0, 0, W, Inches(3.1))
        bg.fill.solid()
        bg.fill.fore_color.rgb = NAVY
        bg.line.fill.background()
        bg.shadow.inherit = False
        self.text(s, "ניתוח כדאיות כלכלית", Inches(1.0), size=40, bold=True,
                  color=WHITE, align=PP_ALIGN.CENTER)
        self.text(s, meta.get("project_name", ""), Inches(1.9), size=24,
                  color=WHITE, align=PP_ALIGN.CENTER)
        details = [v for v in [meta.get("client"), meta.get("location"),
                               meta.get("preparer"), meta.get("date")] if v]
        self.text(s, "  |  ".join(details), Inches(3.5), size=15,
                  color=GREY, align=PP_ALIGN.CENTER)
        self.footer(s)

    def s_project(self):
        s = self.slide("הפרויקט")
        meta = self.m["meta"]
        rows = [[label, meta.get(key) or "—"] for label, key in
                [("מיקום", "location"), ("גוש / חלקה / מגרש", None),
                 ("מצב תכנוני", "plan"), ("מזמין", "client")]]
        rows[1][1] = " / ".join(str(meta.get(k, "—")) for k in ("gush", "helka", "migrash"))
        self.table(s, ["פרט", "תוכן"], rows, Inches(1.8), col_widths=[1, 2.4])
        self.footer(s)

    def s_program(self):
        s = self.slide("תכנית הבנייה והשטחים")
        rows = [[l["label"], "%s מ\"ר" % format(round(l["sqm"]), ","),
                 money(l["cost_per_sqm"]) if l["cost_per_sqm"] else "—",
                 money(l["amount"], short=True)]
                for l in self.m["direct"]["lines"]]
        rows.append(["סה\"כ בנייה ישירה", "", "", money(self.m["direct"]["total"], short=True)])
        self.table(s, ["רכיב", "שטח / כמות", "עלות ליחידה", "סה\"כ"], rows,
                   Inches(1.8), col_widths=[2.2, 1, 1, 1])
        self.footer(s)

    def s_sources_uses(self):
        s = self.slide("מקורות ושימושים")
        m = self.m
        uses = [
            ["קרקע ורכישה", money(m["land"]["total"], short=True)],
            ["בנייה ישירה", money(m["direct"]["total"], short=True)],
            ["עלויות עקיפות", money(m["indirect"]["total"], short=True)],
        ] + ([["מזה: טיפול בדיירים", money(m["tenants"]["total"], short=True)]]
             if (m.get("tenants") or {}).get("lines") else []) + [
            ["בצ\"מ", money(m["contingency"]["amount"], short=True)],
            ["מימון", money(m["finance"]["total"], short=True)],
            ["סה\"כ שימושים", money(m["results"]["total_cost"], short=True)],
        ]
        sources = [
            ["הון עצמי", money(m["finance"]["equity"], short=True)],
            ["אשראי ליווי (שיא)", money(m["results"]["peak_debt"], short=True)],
            ["הכנסות מהשכרה ומימוש (נטו)" if self.kind == "income"
             else "הכנסות ממכירות (נטו)",
             money(m["results"]["total_revenue_net"], short=True)],
        ]
        self.text(s, "שימושים", Inches(1.7), size=18, bold=True, color=NAVY,
                  left=int(W / 2) + Inches(0.2), width=int(W / 2) - Inches(0.8))
        self.text(s, "מקורות", Inches(1.7), size=18, bold=True, color=NAVY,
                  left=Inches(0.6), width=int(W / 2) - Inches(0.8))
        # שתי טבלאות זו לצד זו: שימושים מימין (הקריאה מתחילה שם), מקורות משמאל
        self._half_table(s, ["סעיף", "סכום"], uses, Inches(2.3), right=True)
        self._half_table(s, ["מקור", "סכום"], sources, Inches(2.3), right=False)
        self.footer(s)

    def _half_table(self, slide, headers, rows, top, right=True):
        width = int(W / 2) - Inches(0.8)
        left = (int(W / 2) + Inches(0.2)) if right else Inches(0.6)
        shape = slide.shapes.add_table(len(rows) + 1, 2, left, top, width,
                                       Inches(0.4) * (len(rows) + 1))
        tbl = shape.table
        for c, h in enumerate(reversed(headers)):
            cell = tbl.cell(0, c)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            rtl(p, PP_ALIGN.CENTER)
            r = p.add_run()
            r.text = h
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.name = FONT
            r.font.color.rgb = WHITE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY
        for ri, row in enumerate(rows, start=1):
            for ci, val in enumerate(reversed(row)):
                cell = tbl.cell(ri, ci)
                cell.text = ""
                p = cell.text_frame.paragraphs[0]
                rtl(p, PP_ALIGN.RIGHT if ci == 1 else PP_ALIGN.CENTER)
                r = p.add_run()
                r.text = str(val)
                r.font.size = Pt(13)
                r.font.name = FONT
                r.font.bold = ri == len(rows)

    def s_profit(self):
        """שקף ההחלטה. אם יש שקף אחד שהיזם יזכור, זה הוא."""
        s = self.slide("רווחיות הפרויקט")
        r = self.m["results"]
        ok = r["meets_threshold"]
        self.kpi_row(s, [
            ("רווח יזמי", money(r["profit_before_tax"], short=True), GREEN if ok else RED),
            ("% מהעלויות", percent(r["margin_on_cost"]), GREEN if ok else RED),
            ("IRR שנתי", percent(r["irr_annual"]) if r["irr_annual"] is not None else "—", NAVY),
            ("עלות למ\"ר בנוי" if self.kind == "income" else "עלות למ\"ר מכור",
             money(r["cost_per_sold_sqm"]), NAVY),
        ], top=Inches(1.7))
        self.table(s, ["מדד", "סכום"], [
            ["סך הכנסות (נטו ממע\"מ)", money(r["total_revenue_net"])],
            ["סך עלויות הפרויקט", money(r["total_cost"])],
            ["רווח יזמי לפני מס", money(r["profit_before_tax"])],
            ["רווח לאחר מס", money(r["profit_after_tax"])],
        ], Inches(3.8), col_widths=[2, 1])
        self.text(s, "סף רווח יזמי נדרש: %s — הפרויקט %s"
                  % (percent(r["profit_threshold"], 0),
                     "עומד בסף" if ok else "אינו עומד בסף"),
                  Inches(6.1), size=17, bold=True, color=GREEN if ok else RED)
        self.footer(s)

    def s_sensitivity(self):
        s = self.slide("ניתוח רגישות — מחיר מכירה מול עלות בנייה")
        sen = self.d["sensitivity"]
        headers = ["עלות בנייה ↓ / מחיר →"] + ["%+d%%" % p for p in sen["price_steps"]]
        rows = [["%+d%%" % sen["cost_steps"][i]] + [percent(v) for v in row]
                for i, row in enumerate(sen["grid"])]
        self.table(s, headers, rows, Inches(1.8), size=13)
        th = self.m["results"]["profit_threshold"]
        self.text(s, "כל תא = רווח יזמי כאחוז מהעלויות. סף נדרש: %s. "
                     "תא מתחת לסף מסמן תרחיש שבו הפרויקט מפסיק להיות כדאי."
                  % percent(th, 0), Inches(5.4), size=13, color=GREY)
        self.footer(s)

    def s_rates(self):
        s = self.slide("רגישות לריבית")
        rows = [["%+.1f נק' אחוז" % sc["delta_pp"], percent(sc["rate"]),
                 money(sc["interest"], short=True), money(sc["profit"], short=True),
                 percent(sc["metric"])]
                for sc in self.d["rate_scenarios"]]
        self.table(s, ["שינוי בריבית", "ריבית שנתית", "עלות ריבית", "רווח יזמי", "% מהעלויות"],
                   rows, Inches(1.8))
        self.footer(s)

    def s_cashflow(self):
        s = self.slide("תזרים ומימון")
        f = self.m["finance"]
        self.kpi_row(s, [
            ("הון עצמי", money(f["equity"], short=True), NAVY),
            ("שיא ניצול אשראי", money(f["peak_debt"], short=True), AMBER),
            ("סה\"כ עלויות מימון", money(f["total"], short=True), NAVY),
        ], top=Inches(1.7))
        peak = max((r["debt_balance"] for r in self.m["cashflow"]), default=0)
        rows = []
        for r in self.m["cashflow"]:
            bar = "█" * int(round(12 * r["debt_balance"] / peak)) if peak else ""
            rows.append([r["quarter"], money(r["outflow"], short=True),
                         money(r["revenue"], short=True),
                         money(r["debt_balance"], short=True), bar])
        self.table(s, ["רבעון", "הוצאות", "הכנסות", "יתרת אשראי", "עקומת חוב"],
                   rows, Inches(3.7), size=11, col_widths=[1, 1.1, 1.1, 1.2, 1.6])
        self.footer(s)

    def s_flags(self):
        s = self.slide("דגלים ונקודות לבחינה")
        flags = self.d["flags"]
        if not flags:
            self.bullets(s, ["המודל לא זיהה חריגה מהספים שהוגדרו.",
                             "אין בכך אישור לפרויקט — הדגלים נגזרים מהמספרים בלבד."],
                         Inches(2.0))
        else:
            items = ["%s  %s — %s" % ("🔴" if f.get("level") == "red" else "🟡",
                                      f.get("title", ""), f.get("text", ""))
                     for f in flags[:6]]
            self.bullets(s, items, Inches(1.7), size=14)
        self.footer(s)

    def s_recommendation(self):
        s = self.slide("המלצה")
        r = self.m["results"]
        ok = r["meets_threshold"]
        reds = [f for f in self.d["flags"] if f.get("level") == "red"]
        if ok and not reds:
            headline, color = "הפרויקט כדאי כלכלית", GREEN
            body = ("הרווח היזמי %s (%s מהעלויות) מעל הסף הנדרש, ולא אותרו דגלים "
                    "אדומים. מומלץ להתקדם, בכפוף לאימות ההנחות מול אסמכתאות מעודכנות."
                    % (money(r["profit_before_tax"]), percent(r["margin_on_cost"])))
        elif ok and reds:
            headline, color = "הפרויקט כדאי — בכפוף לטיפול בדגלים", AMBER
            body = ("הרווח היזמי עומד בסף (%s מהעלויות), אך אותרו %d דגלים אדומים "
                    "שיש להסיר לפני החלטת השקעה: %s."
                    % (percent(r["margin_on_cost"]), len(reds),
                       "; ".join(f["title"] for f in reds)))
        elif r.get("threshold_basis") == "spread":
            im = self.m.get("income_metrics") or {}
            headline, color = "הפרויקט אינו עומד במבחן המרווח", RED
            body = ("התשואה על העלות %s מול שיעור היוון ביציאה %s — מרווח של %d נקודות "
                    "בסיס בלבד, מול נורמה של 150–200. הנכס שווה כמעט בדיוק את עלות "
                    "הקמתו, ותנודה קטנה בשכירות או בהיוון מוחקת את הרווח."
                    % (percent(im.get("yield_on_cost"), 2), percent(im.get("exit_cap_rate"), 2),
                       round(im.get("spread_bps") or 0)))
        else:
            headline, color = "הפרויקט אינו עומד בסף הכדאיות", RED
            body = ("הרווח היזמי %s מהעלויות, מתחת לסף %s. נדרש שינוי מהותי — מחיר "
                    "הקרקע, תמהיל השטחים, מחירי המכירה או מבנה המימון — בטרם החלטת "
                    "השקעה." % (percent(r["margin_on_cost"]), percent(r["profit_threshold"], 0)))
        self.text(s, headline, Inches(2.0), size=30, bold=True, color=color)
        self.text(s, body, Inches(3.0), size=17)
        self.text(s, "הצעדים הבאים:", Inches(4.3), size=17, bold=True, color=NAVY)
        steps = ["אימות דמי השכירות ושיעור ההיוון מול עסקאות השוואה לנכסים דומים.",
                 "בחינת חוזה עוגן — הוא מרחיב את שיעור הליווי ומוזיל את ההון הנדרש.",
                 "אימות זכויות הבנייה בתב\"ע מול השטח הבנוי המתוכנן."] if self.kind == "income" else [
                 "אימות מחירי המכירה מול עסקאות השוואה עדכניות באזור.",
                 "קבלת אומדן היטל השבחה מהוועדה המקומית.",
                 "בחינת תנאי הליווי מול הבנק המלווה."]
        self.bullets(s, steps, Inches(4.9), size=15)
        self.footer(s)

    def build(self, path):
        self.s_cover()
        self.s_project()
        self.s_program()
        self.s_sources_uses()
        self.s_profit()
        self.s_cashflow()
        self.s_sensitivity()
        self.s_rates()
        self.s_flags()
        self.s_recommendation()
        self.prs.save(path)
        return path


def main():
    ap = argparse.ArgumentParser(description="בונה מצגת PPTX ליזם")
    ap.add_argument("project")
    ap.add_argument("-o", "--output", default="מצגת ליזם.pptx")
    args = ap.parse_args()
    data = full_output(load_project(args.project))
    path = Deck(data).build(args.output)
    print("נוצר: %s  (%d שקפים)" % (path, len(Presentation(path).slides.__iter__.__self__._sldIdLst)))
    return 0


if __name__ == "__main__":
    sys.exit(main())

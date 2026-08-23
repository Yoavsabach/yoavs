#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_report.py — בונה את הדוח הנרטיבי כקובץ Word בעברית RTL מלא.

כללי ה-RTL כאן אינם המצאה מקומית: הם מיושמים לפי hebrew-doc-studio
(references/rtl-bidi-cheatsheet.md), ושלושת התנאים שחייבים להתקיים יחד הם —
  1. בסיס פסקה RTL (``w:bidi``) במיקום הנכון בסכמת ``w:pPr``.
  2. שורה מעורבת נשארת ב-run אחד. Word מריץ את אלגוריתם ה-bidi בעצמו; פיצול
     ידני לפי כתב שובר רווחים, פיסוק, סוגריים ותאריכים.
  3. גופן וגודל גם על ``w:cs``/``w:szCs`` — עברית היא "כתב מורכב" ב-Word,
     ו-``w:ascii``/``w:sz`` לבדם לא חלים עליה.
מלכודת נוספת שמיושמת כאן: תחת ``w:bidi``, ``w:jc="right"`` מתפרש כקצה הלוגי
ומרונדר **שמאלה**. ליישור-ימין פשוט לא מגדירים ``w:jc`` כלל.

שימוש:
    python3 build_report.py project.json -o "דוח אפס.docx" [--mode takan21|short]
"""

from __future__ import annotations

import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from model import DISCLAIMER, full_output, load_project, num  # noqa: E402

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor
except ImportError:  # pragma: no cover
    print("חסרה תלות. התקן:  pip install python-docx", file=sys.stderr)
    raise

FONT = "David"          # קיים כמעט בכל Windows עם Office — ראה hebrew-doc-studio
SIZE = 12
NAVY = "1F3864"

# סדר הילדים החוקי של CT_PPr. w:bidi חייב לבוא לפני w:jc, אחרת Word
# ו-Google Docs מתעלמים מה-bidi והפסקה נשפכת שמאלה.
_PPR_ORDER = [
    'pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl',
    'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens',
    'kinsoku', 'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE',
    'autoSpaceDN', 'bidi', 'adjustRightInd', 'snapToGrid', 'spacing', 'ind',
    'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc', 'textDirection',
    'textAlignment', 'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr',
    'sectPr', 'pPrChange',
]


def _insert_ordered(pPr, element):
    tag = element.tag.split('}')[-1]
    idx = _PPR_ORDER.index(tag)
    for child in pPr:
        ctag = child.tag.split('}')[-1]
        if ctag in _PPR_ORDER and _PPR_ORDER.index(ctag) > idx:
            child.addprevious(element)
            return
    pPr.append(element)


def rtl_para(paragraph, align="right"):
    pPr = paragraph._p.get_or_add_pPr()
    _insert_ordered(pPr, pPr.makeelement(qn('w:bidi'), {}))
    # ליישור-ימין לא מגדירים jc — קצה ההובלה של פסקת bidi הוא ימין בכל renderer.
    if align == "center":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return paragraph


def rtl_run(paragraph, text, bold=False, italic=False, size=SIZE, color=None, font=FONT):
    """כותב run יחיד עם כל דגלי הכתב המורכב. השורה כולה ב-run אחד במכוון."""
    run = paragraph.add_run(text)
    rPr = run._r.get_or_add_rPr()
    rPr.append(rPr.makeelement(qn('w:rFonts'), {
        qn('w:ascii'): font, qn('w:hAnsi'): font, qn('w:cs'): font}))
    if bold:
        rPr.append(rPr.makeelement(qn('w:b'), {}))
        rPr.append(rPr.makeelement(qn('w:bCs'), {}))
    if italic:
        rPr.append(rPr.makeelement(qn('w:i'), {}))
        rPr.append(rPr.makeelement(qn('w:iCs'), {}))
    if color:
        rPr.append(rPr.makeelement(qn('w:color'), {qn('w:val'): color}))
    rPr.append(rPr.makeelement(qn('w:sz'), {qn('w:val'): str(int(size * 2))}))
    rPr.append(rPr.makeelement(qn('w:szCs'), {qn('w:val'): str(int(size * 2))}))
    rPr.append(rPr.makeelement(qn('w:rtl'), {}))
    return run


def money(x):
    try:
        return "%s ₪" % format(int(round(float(x))), ",")
    except (TypeError, ValueError):
        return "—"


def percent(x, digits=1):
    try:
        return ("%." + str(digits) + "f%%") % (float(x) * 100)
    except (TypeError, ValueError):
        return "—"


class Report:
    def __init__(self, data, mode="takan21"):
        self.d = data
        self.m = data["model"]
        self.mode = mode
        self.doc = Document()
        self._setup()

    def _setup(self):
        sec = self.doc.sections[0]
        sec.left_margin = sec.right_margin = Mm(25)
        sec.top_margin = sec.bottom_margin = Mm(22)
        # כיוון הסקשן — בלעדיו העמוד עצמו LTR והכותרות העליונות מתהפכות.
        sec._sectPr.append(sec._sectPr.makeelement(qn('w:bidi'), {}))
        style = self.doc.styles["Normal"]
        style.font.name = FONT
        style.font.size = Pt(SIZE)
        rpr = style.element.get_or_add_rPr()
        rpr.append(rpr.makeelement(qn('w:rFonts'), {
            qn('w:ascii'): FONT, qn('w:hAnsi'): FONT, qn('w:cs'): FONT}))

    # -- אבני בניין --------------------------------------------------------

    def h(self, text, level=1):
        sizes = {0: 20, 1: 15, 2: 13, 3: 12}
        p = self.doc.add_paragraph()
        rtl_para(p, "center" if level == 0 else "right")
        rtl_run(p, text, bold=True, size=sizes.get(level, 12), color=NAVY)
        p.paragraph_format.space_before = Pt(14 if level <= 1 else 10)
        p.paragraph_format.space_after = Pt(6)
        return p

    def para(self, text, bold=False, size=SIZE, align="justify"):
        p = self.doc.add_paragraph()
        rtl_para(p, align)
        rtl_run(p, text, bold=bold, size=size)
        p.paragraph_format.line_spacing = 1.4
        p.paragraph_format.space_after = Pt(6)
        return p

    def bullets(self, items):
        for item in items:
            p = self.doc.add_paragraph(style="List Bullet")
            rtl_para(p)
            rtl_run(p, item)
            p.paragraph_format.line_spacing = 1.35

    def table(self, headers, rows, widths=None):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        # bidiVisual — בלעדיו העמודה הראשונה נופלת משמאל גם במסמך RTL.
        tblPr = t._tbl.tblPr
        tblPr.append(tblPr.makeelement(qn('w:bidiVisual'), {}))
        for i, htext in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            rtl_para(p, "center")
            rtl_run(p, htext, bold=True, size=11, color="FFFFFF")
            shd = cell._tc.get_or_add_tcPr().makeelement(
                qn('w:shd'), {qn('w:val'): 'clear', qn('w:fill'): NAVY})
            cell._tc.get_or_add_tcPr().append(shd)
        for row in rows:
            cells = t.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = ""
                p = cells[i].paragraphs[0]
                rtl_para(p, "right" if i == 0 else "center")
                rtl_run(p, str(value), size=11)
        if widths:
            for i, w in enumerate(widths):
                for row in t.rows:
                    row.cells[i].width = Mm(w)
        self.doc.add_paragraph()
        return t

    def page_break(self):
        self.doc.add_page_break()

    # -- פרקי הדוח ---------------------------------------------------------

    def cover(self):
        meta = self.m["meta"]
        for _ in range(4):
            self.doc.add_paragraph()
        self.h("ניתוח כדאיות כלכלית — דוח אפס", 0)
        p = self.doc.add_paragraph()
        rtl_para(p, "center")
        rtl_run(p, meta.get("project_name", ""), bold=True, size=16)
        for label, key in [("מזמין הדוח", "client"), ("מיקום", "location"),
                           ("עורך הדוח", "preparer"), ("תאריך", "date")]:
            if meta.get(key):
                p = self.doc.add_paragraph()
                rtl_para(p, "center")
                rtl_run(p, "%s: %s" % (label, meta[key]), size=12)
        ident = " | ".join("%s %s" % (k, meta[v]) for k, v in
                           [("גוש", "gush"), ("חלקה", "helka"), ("מגרש", "migrash")]
                           if meta.get(v))
        if ident:
            p = self.doc.add_paragraph()
            rtl_para(p, "center")
            rtl_run(p, ident, size=12, bold=True)
        for _ in range(3):
            self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        rtl_para(p, "center")
        rtl_run(p, DISCLAIMER, italic=True, size=10, color="7F7F7F")
        self.page_break()

    def executive_summary(self):
        """תמצית מנהלים בעמוד הראשון.

        היזם קורא את העמוד הזה ומחליט; שאר הדוח הוא ההנמקה. לכן המסקנה נאמרת
        כאן במפורש — כולל מסקנה שלילית — ולא נגררת לסוף.
        """
        r = self.m["results"]
        self.h("תמצית מנהלים", 1)
        meta = self.m["meta"]
        self.para("דוח זה בוחן את הכדאיות הכלכלית של %s. הבחינה נערכה על בסיס הנתונים "
                  "התכנוניים, עלויות הבנייה ומחירי המכירה המפורטים בפרקים שלהלן."
                  % (meta.get("project_name") or "הפרויקט"))

        self.table(
            ["מדד", "תוצאה"],
            [
                ["סך עלויות הפרויקט", money(r["total_cost"])],
                ["סך הכנסות צפויות (נטו ממע\"מ)", money(r["total_revenue_net"])],
                ["רווח יזמי לפני מס", money(r["profit_before_tax"])],
                ["רווח יזמי — % מהעלויות", percent(r["margin_on_cost"])],
                ["רווח יזמי — % מההכנסות", percent(r["margin_on_revenue"])],
                ["IRR שנתי על ההון העצמי",
                 percent(r["irr_annual"]) if r["irr_annual"] is not None else "לא ניתן לחישוב"],
                ["עלות למ\"ר מכור", money(r["cost_per_sold_sqm"])],
                ["שיא ניצול אשראי ליווי", money(r["peak_debt"])],
            ], widths=[90, 60])

        th = r["profit_threshold"]
        if r["meets_threshold"]:
            verdict = ("הפרויקט מציג רווח יזמי של %s, המהווה %s מסך העלויות — מעל סף "
                       "ה-%s הנהוג בליווי בנקאי. בכפוף לאימות ההנחות ולטיפול בדגלים "
                       "המפורטים בפרק הדגלים, הפרויקט נמצא כדאי כלכלית."
                       % (money(r["profit_before_tax"]), percent(r["margin_on_cost"]),
                          percent(th, 0)))
        else:
            verdict = ("הפרויקט מציג רווח יזמי של %s, המהווה %s מסך העלויות — מתחת לסף "
                       "ה-%s הנהוג בליווי בנקאי. במתכונת הנוכחית הפרויקט אינו עומד "
                       "בסף הכדאיות, ונדרש שינוי מהותי בפרמטרים (מחיר הקרקע, תמהיל "
                       "השטחים, מחירי המכירה או מבנה המימון) בטרם קבלת החלטת השקעה."
                       % (money(r["profit_before_tax"]), percent(r["margin_on_cost"]),
                          percent(th, 0)))
        self.h("מסקנה", 2)
        self.para(verdict, bold=True)

        reds = [f for f in self.d["flags"] if f.get("level") == "red"]
        if reds:
            self.h("דגלים אדומים מרכזיים", 2)
            self.bullets(["%s — %s" % (f["title"], f["text"]) for f in reds])
        self.page_break()

    def identification(self):
        meta = self.m["meta"]
        self.h("1. זיהוי המקרקעין ומטרת הדוח", 1)
        self.para("מטרת הדוח היא בחינת הכדאיות הכלכלית של הפרויקט עבור היזם, לצורך "
                  "קבלת החלטת השקעה ובחינת התכנות מול גורם מממן.")
        rows = [[label, meta.get(key, "—") or "—"] for label, key in
                [("שם הפרויקט", "project_name"), ("מיקום", "location"), ("גוש", "gush"),
                 ("חלקה", "helka"), ("מגרש", "migrash"), ("מצב תכנוני", "plan"),
                 ("מזמין הדוח", "client"), ("תאריך עריכה", "date")]]
        self.table(["פרט", "תוכן"], rows, widths=[55, 95])

    def planning(self):
        self.h("2. המצב התכנוני וזכויות הבנייה", 1)
        meta = self.m["meta"]
        if meta.get("plan"):
            self.para("המצב התכנוני החל על המקרקעין: %s." % meta["plan"])
        self.para("להלן פירוט השטחים שנלקחו בחשבון בבניית המודל, לפי שימוש:")
        rows = [[line["label"], "%s מ\"ר" % format(round(line["sqm"]), ","),
                 money(line["cost_per_sqm"]) if line["cost_per_sqm"] else "—",
                 line.get("source") or "—"]
                for line in self.m["direct"]["lines"]]
        self.table(["רכיב", "שטח / כמות", "עלות ליחידה", "מקור"], rows, widths=[52, 30, 32, 36])
        self.para("אומדני היטל ההשבחה ודמי ההיתר, ככל שנכללו, הם אומדנים בלבד. "
                  "קביעה מחייבת נעשית בשומת הוועדה המקומית או ברמ\"י, וניתנת להשגה "
                  "בפני שמאי מכריע.")

    def costs(self):
        self.h("3. אומדן העלויות", 1)
        self.h("3.1 קרקע ורכישה", 2)
        self.table(["סעיף", "סכום"],
                   [[l["label"], money(l["amount"])] for l in self.m["land"]["lines"]] +
                   [["סה\"כ קרקע ורכישה", money(self.m["land"]["total"])]], widths=[95, 55])

        self.h("3.2 עלויות בנייה ישירות", 2)
        self.table(["רכיב", "שטח / כמות", "עלות ליחידה", "סה\"כ"],
                   [[l["label"], format(round(l["sqm"]), ","),
                     money(l["cost_per_sqm"]) if l["cost_per_sqm"] else "—",
                     money(l["amount"])] for l in self.m["direct"]["lines"]] +
                   [["סה\"כ עלויות ישירות", "", "", money(self.m["direct"]["total"])]],
                   widths=[52, 28, 32, 38])

        self.h("3.3 עלויות עקיפות", 2)
        self.table(["סעיף", "שיעור", "סכום"],
                   [[l["label"], percent(l["pct"]) if l["pct"] else "—", money(l["amount"])]
                    for l in self.m["indirect"]["lines"]] +
                   [["סה\"כ עלויות עקיפות", "", money(self.m["indirect"]["total"])]],
                   widths=[80, 30, 40])

        self.h("3.4 בלתי צפוי מראש (בצ\"מ)", 2)
        c = self.m["contingency"]
        self.para("בצ\"מ נקבע בשיעור %s מבסיס של %s, ובסך %s. הטווח המקובל בליווי "
                  "בנקאי הוא 5%%–10%%; שיעור נמוך מכך נבחן על ידי הבנק המלווה."
                  % (percent(c["pct"]), money(c["base"]), money(c["amount"])))

    def finance(self):
        self.h("4. מבנה המימון", 1)
        f = self.m["finance"]
        self.para("מקורות המימון: הון עצמי בסך %s ואשראי ליווי. שיא ניצול האשראי לאורך "
                  "חיי הפרויקט עומד על %s, וזו מסגרת האשראי הנדרשת. הריבית חושבה "
                  "בשיעור שנתי של %s על היתרה המנוצלת בפועל בכל רבעון."
                  % (money(f["equity"]), money(f["peak_debt"]), percent(f["annual_rate"])))
        self.table(["סעיף", "סכום"],
                   [[l["label"], money(l["amount"])] for l in f["lines"]] +
                   [["סה\"כ עלויות מימון", money(f["total"])]], widths=[95, 55])

    def revenue(self):
        self.h("5. אומדן ההכנסות", 1)
        rev = self.m["revenue"]
        self.para("מחירי המכירה הוזנו %s מע\"מ. הרווח היזמי נמדד על ההכנסה נטו, שכן "
                  "המע\"מ מועבר לרשות המסים ואינו מהווה הכנסה של היזם."
                  % ("כוללי" if rev["prices_include_vat"] else "ללא"))
        self.table(["רכיב", "יחידות", "שטח ממוצע", "מחיר למ\"ר", "סה\"כ ברוטו", "נטו"],
                   [[l["label"], format(int(l["units"]), ","),
                     format(round(l["avg_sqm"]), ","), money(l["price_per_sqm"]),
                     money(l["gross"]), money(l["net"])] for l in rev["lines"]] +
                   [["סה\"כ", "", "", "", money(rev["gross"]), money(rev["net"])]],
                   widths=[36, 18, 22, 26, 26, 26])

    def cashflow(self):
        self.h("6. תזרים המזומנים", 1)
        self.para("התזרים נפרס על פני %d רבעונים. ההון העצמי נספג ראשון, והריבית נצברת "
                  "על היתרה שמעבר לו בלבד." % len(self.m["cashflow"]))
        rows = [[r["quarter"], money(r["outflow"]), money(r["revenue"]),
                 money(r["interest"]), money(r["debt_balance"])]
                for r in self.m["cashflow"]]
        self.table(["רבעון", "הוצאות", "הכנסות", "ריבית", "יתרת אשראי"], rows,
                   widths=[24, 32, 32, 28, 34])

    def results(self):
        self.h("7. תוצאות המודל", 1)
        r = self.m["results"]
        self.table(["מדד", "תוצאה"], [
            ["סך עלויות הפרויקט", money(r["total_cost"])],
            ["סך הכנסות נטו", money(r["total_revenue_net"])],
            ["רווח יזמי לפני מס", money(r["profit_before_tax"])],
            ["רווח יזמי — % מהעלויות", percent(r["margin_on_cost"])],
            ["רווח יזמי — % מההכנסות", percent(r["margin_on_revenue"])],
            ["מס על הרווח", money(r["profit_tax"])],
            ["רווח לאחר מס", money(r["profit_after_tax"])],
            ["IRR שנתי", percent(r["irr_annual"]) if r["irr_annual"] is not None else "לא ניתן לחישוב"],
            ["תשואה על ההון (ROC)", percent(r["roc"])],
            ["עלות למ\"ר מכור", money(r["cost_per_sold_sqm"])],
        ], widths=[90, 60])

    def sensitivity(self):
        self.h("8. ניתוחי רגישות", 1)
        s = self.d["sensitivity"]
        self.para("הטבלה מציגה את הרווח היזמי כאחוז מהעלויות בשילובים של שינוי במחיר "
                  "המכירה ובעלות הבנייה. תא שלילי או נמוך מהסף מסמן תרחיש שבו הפרויקט "
                  "מפסיק להיות כדאי.")
        headers = ["עלות בנייה \\ מחיר מכירה"] + ["%+d%%" % p for p in s["price_steps"]]
        rows = [["%+d%%" % s["cost_steps"][i]] + [percent(v) for v in row]
                for i, row in enumerate(s["grid"])]
        self.table(headers, rows)

        self.h("8.1 תרחישי ריבית", 2)
        self.table(["שינוי בריבית", "ריבית שנתית", "עלות ריבית", "רווח יזמי", "% מהעלויות"],
                   [["%+.1f נק' אחוז" % sc["delta_pp"], percent(sc["rate"]),
                     money(sc["interest"]), money(sc["profit"]), percent(sc["metric"])]
                    for sc in self.d["rate_scenarios"]])

    def flags(self):
        self.h("9. דגלים ונקודות לבחינה", 1)
        flags = self.d["flags"]
        if not flags:
            self.para("המודל לא זיהה חריגה מהספים שהוגדרו. אין בכך משום אישור לפרויקט — "
                      "הדגלים נגזרים מהמספרים בלבד ואינם מחליפים בדיקה תכנונית ומשפטית.")
        for f in flags:
            self.h("%s %s" % ("🔴" if f.get("level") == "red" else "🟡", f.get("title", "")), 3)
            self.para(f.get("text", ""))
        devs = self.d["deviations"]
        if devs:
            self.h("9.1 סטיות מאסמכתאות חיצוניות", 2)
            self.para("הקלטים הבאים סוטים ביותר מ-10%% מהאסמכתה החיצונית שאותרה. הקלט "
                      "לא שונה — ההחלטה על המספר היא של היזם — אך הפער מדווח במפורש:")
            self.bullets([d["text"] for d in devs])

    def methodology(self):
        self.h("10. מתודולוגיה, מקורות ומגבלות", 1)
        self.para("המודל נבנה כמודל תזרימי רבעוני. העלויות והכנסות נפרסות על ציר הזמן, "
                  "עלות המימון מחושבת על היתרה המנוצלת בפועל בכל רבעון, והרווח היזמי "
                  "נמדד הן כאחוז מהעלויות והן כאחוז מההכנסות.")
        self.h("10.1 מקורות הנתונים", 2)
        sources = []
        for group in ("land", "direct", "indirect"):
            for line in self.m[group]["lines"]:
                if line.get("source"):
                    sources.append("%s — %s" % (line["label"], line["source"]))
        for line in self.m["revenue"]["lines"]:
            if line.get("source"):
                sources.append("%s — %s" % (line["label"], line["source"]))
        self.bullets(sources or ["לא צוינו מקורות פרטניים בקלט. יש להשלים לפני הצגה לגורם מממן."])
        self.h("10.2 מגבלות", 2)
        self.bullets([
            "הדוח מבוסס על נתונים שנמסרו על ידי המזמין ועל אסמכתאות חיצוניות שאותרו במועד עריכתו.",
            "אומדני היטל השבחה ודמי היתר הם אומדנים; הקביעה המחייבת היא של הוועדה המקומית או רמ\"י.",
            "שיעורי המס ומחירי התשומות משתנים; יש לאמת אותם מחדש בכל עדכון של הדוח.",
            "המודל אינו כולל בחינה משפטית של הזכויות במקרקעין ואינו מהווה בדיקת נאותות.",
        ])
        p = self.doc.add_paragraph()
        rtl_para(p)
        rtl_run(p, DISCLAIMER, bold=True, size=11, color="C00000")

    # -- הרכבה -------------------------------------------------------------

    def build(self, path):
        self.cover()
        self.executive_summary()
        if self.mode == "short":
            # הגרסה היזמית המקוצרת: מה שצריך להחלטת Go/No-Go, בלי מנגנון התקן.
            self.costs()
            self.revenue()
            self.results()
            self.sensitivity()
            self.flags()
            self.methodology()
        else:
            self.identification()
            self.planning()
            self.costs()
            self.finance()
            self.revenue()
            self.cashflow()
            self.results()
            self.sensitivity()
            self.flags()
            self.methodology()
        self.doc.save(path)
        return path


def main():
    ap = argparse.ArgumentParser(description="בונה את הדוח הנרטיבי כקובץ Word")
    ap.add_argument("project")
    ap.add_argument("-o", "--output", default="דוח אפס.docx")
    ap.add_argument("--mode", choices=["takan21", "short"], default=None,
                    help="ברירת מחדל: לפי meta.mode שב-project.json")
    args = ap.parse_args()
    project = load_project(args.project)
    mode = args.mode or (project.get("meta", {}) or {}).get("mode") or "takan21"
    data = full_output(project)
    path = Report(data, mode=mode).build(args.output)
    print("נוצר: %s  (מצב: %s)" % (path, mode))
    return 0


if __name__ == "__main__":
    sys.exit(main())
